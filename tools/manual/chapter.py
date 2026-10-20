"""chapter.py TEXT.md [...]: writes chapters of the manual from their text (docs/Manual_Plan.md, M3).

Each file in docs/manual/text/ is one chapter. Its first line, "# Title", names the chapter's Heading 1 in
docs/manual/Colour_Maximite_2_User_Manual.docx. Everything between that heading and the next Heading 1 is replaced by
the file's text, in the PicoMite manual's own formats, so running it again after an edit is safe:

    ## Heading           Heading 2 (### for Heading 3)
    a paragraph         Normal; lines join up to a blank line. **bold**, *italic* and `code` (Courier New) inline
    - item              a bullet (the manual's list: numbering 5); a line indented by two spaces goes on with it
    ```                 code, a line a paragraph: Courier New, indented 540, no spacing (**bold** for what is typed)
    | a | b |           a table: the first row is the header (shaded); "{widths: 10,40}" on the line before sets the
                        columns in parts of the text width
    ![](pdf:P:N){width=12}   a picture: image xref N of page P of the b14 PDF (or a path), width in cm, centred
    >> a<TAB>b          an indented line with tab-separated columns (a list such as the shortcut keys)
    !TODO text          a highlighted [MANUAL TODO] note: something still to check
    !NOTE text          a note in italics (a paragraph of its own)

The paragraph after a code block, a table or a picture gets the template's 120 before it."""
import io, os, re, sys
import docx
import fitz
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
MANUAL = os.path.join(ROOT, 'docs', 'manual', 'Colour_Maximite_2_User_Manual.docx')
B14 = r'D:\Dropbox\VSCode-STM32\Colour_Maximite_2_User_Manual.pdf'
TODO = '[MANUAL TODO] '
TEXT_WIDTH = 9866                                       # twips: A4 less the 18 mm margins
BULLETS = '5'                                           # the manual's bullet list (numbering.xml)


def ptext(el):
    return ''.join(t.text or '' for t in el.iter(qn('w:t'))).strip()


class Writer:
    def __init__(self, d):
        self.d = d
        self.out = []                                   # the new elements, in order
        self.after_block = False
        self.pdf = None

    # paragraphs are made at the end of the body by python-docx, then moved into place
    def para(self, style=None):
        p = self.d.add_paragraph(style=style)
        self.out.append(p._p)
        if self.after_block and style is None:
            p.paragraph_format.space_before = Pt(6)
        self.after_block = False
        return p

    def runs(self, p, text, code=False):
        for part in re.split(r'(\*\*.+?\*\*|`[^`]+`|(?<![*\w])\*[^*\s][^*]*\*(?![*\w]))', text):
            if not part:
                continue
            if part.startswith('**') and part.endswith('**'):
                r = p.add_run(part[2:-2])
                r.bold = True
            elif part.startswith('`') and part.endswith('`'):
                r = p.add_run(part[1:-1])
                r.font.name = 'Courier New'
            elif part.startswith('*') and part.endswith('*') and len(part) > 2:
                r = p.add_run(part[1:-1])
                r.italic = True
            else:
                r = p.add_run(part)
            if code:
                r.font.name = 'Courier New'

    def heading(self, level, text):
        p = self.para('Heading %d' % level)
        p.add_run(text)

    def paragraph(self, text):
        self.runs(self.para(), text)

    def bullet(self, text):
        p = self.para()
        ppr = p._p.get_or_add_pPr()
        numpr = OxmlElement('w:numPr')
        ilvl = OxmlElement('w:ilvl')
        ilvl.set(qn('w:val'), '0')
        num = OxmlElement('w:numId')
        num.set(qn('w:val'), BULLETS)
        numpr.append(ilvl)
        numpr.append(num)
        ppr.append(numpr)
        p.paragraph_format.left_indent = Cm(540 / 567)
        p.paragraph_format.first_line_indent = -Cm(360 / 567)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        self.runs(p, text)

    def tabline(self, text):
        """an indented line whose columns are separated by tabs (a list such as the shortcut keys)"""
        p = self.para()
        p.paragraph_format.left_indent = Cm(4.5)                   # (a hanging indent: the first column at 1.27 cm,
        p.paragraph_format.first_line_indent = Cm(720 / 567 - 4.5) # the rest at 4.5 cm, a long description wrapping
        p.paragraph_format.space_before = Pt(2)                    # under itself)
        p.paragraph_format.space_after = Pt(2)
        for stop in (4.5, 8.0, 11.5):
            p.paragraph_format.tab_stops.add_tab_stop(Cm(stop))
        for n, col in enumerate(text.split('\t')):
            if n:
                p.add_run('\t')
            self.runs(p, col)

    def code(self, lines):
        for n, line in enumerate(lines):
            p = self.para()
            p.paragraph_format.left_indent = Cm(540 / 567)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6) if n == len(lines) - 1 else Pt(0)
            self.runs(p, line if line else ' ', code=True)
        self.after_block = True

    def note(self, text, todo=True):
        p = self.para()
        r = p.add_run((TODO if todo else '') + text)
        if todo:
            hl = OxmlElement('w:highlight')
            hl.set(qn('w:val'), 'yellow')
            r._r.get_or_add_rPr().append(hl)
        else:
            r.italic = True

    def picture(self, src, width):
        if src.startswith('clip:'):                    # a drawing in the PDF (vectors, not an image): its page area
            _, page, box = src.split(':')
            if self.pdf is None:
                self.pdf = fitz.open(B14)
            rect = fitz.Rect(*[float(v) for v in box.split(',')])
            blob = io.BytesIO(self.pdf[int(page) - 1].get_pixmap(dpi=220, clip=rect).tobytes('png'))
        elif src.startswith('pdf:'):
            _, page, xref = src.split(':')
            if self.pdf is None:
                self.pdf = fitz.open(B14)
            pix = fitz.Pixmap(self.pdf, int(xref))      # (re-encoded: python-docx does not read every JPEG the PDF
            if pix.n - pix.alpha >= 4 or pix.colorspace is None or pix.colorspace.n != 3:   # holds, a CMYK or grey
                pix = fitz.Pixmap(fitz.csRGB, pix)                                          # one among them)
            if pix.alpha:
                blob = io.BytesIO(pix.tobytes('png'))
            else:
                blob = io.BytesIO(pix.tobytes('jpg', jpg_quality=90))
        else:
            blob = os.path.join(ROOT, src)
        p = self.para()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(blob, width=Cm(width))
        self.after_block = True

    def table(self, rows, widths):
        n = len(rows[0])
        widths = widths or [1] * n
        tw = [int(TEXT_WIDTH * w / sum(widths)) for w in widths]
        t = self.d.add_table(rows=len(rows), cols=n)
        t.style = self.d.styles['Table Grid'] if 'Table Grid' in [s.name for s in self.d.styles] else None
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r, row in enumerate(rows):
            for c, cell in enumerate(row):
                tc = t.cell(r, c)
                tc.width = docx.shared.Twips(tw[c])
                p = tc.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                self.runs(p, cell.strip())
                if r == 0:
                    for run in p.runs:
                        run.bold = True
                    shd = OxmlElement('w:shd')
                    shd.set(qn('w:val'), 'clear')
                    shd.set(qn('w:color'), 'auto')
                    shd.set(qn('w:fill'), 'D9D9D9')
                    tc._tc.get_or_add_tcPr().append(shd)
        self.out.append(t._tbl)
        self.after_block = True


def parse(path, w):
    lines = open(path, encoding='utf-8').read().split('\n')
    title = lines[0][2:].strip() if lines[0].startswith('# ') else None
    if not title:
        sys.exit('%s: the first line must be "# Chapter heading"' % path)
    i, para, widths = 1, [], None

    def flush():
        if para:
            w.paragraph(' '.join(s.strip() for s in para))
            para.clear()

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if line.startswith('```'):
            flush()
            j = i + 1
            block = []
            while j < len(lines) and not lines[j].startswith('```'):
                block.append(lines[j].rstrip())
                j += 1
            w.code(block)
            i = j + 1
            continue
        if s.startswith('{widths:'):
            flush()
            widths = [float(x) for x in s[8:-1].split(',')]
        elif s.startswith('|'):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [c for c in lines[i].strip().strip('|').split('|')]
                if not all(re.match(r'^\s*:?-+:?\s*$', c) for c in cells):
                    rows.append([c.replace('\\|', '|') for c in cells])
                i += 1
            w.table(rows, widths)
            widths = None
            continue
        elif s.startswith('### '):
            flush()
            w.heading(3, s[4:])
        elif s.startswith('## '):
            flush()
            w.heading(2, s[3:])
        elif s.startswith('- '):
            flush()
            item = [s[2:]]
            while i + 1 < len(lines) and lines[i + 1].startswith('  ') and lines[i + 1].strip() and \
                    not lines[i + 1].strip().startswith('- '):
                i += 1
                item.append(lines[i].strip())
            w.bullet(' '.join(item))
        elif s.startswith('!['):
            flush()
            m = re.match(r'!\[[^\]]*\]\(([^)]+)\)(?:\{width=([\d.]+)\})?', s)
            w.picture(m.group(1), float(m.group(2) or 15))
        elif line.startswith('>> '):
            flush()
            w.tabline(line[3:].rstrip())
        elif s.startswith('!TODO '):
            flush()
            w.note(s[6:])
        elif s.startswith('!NOTE '):
            flush()
            w.note(s[6:], todo=False)
        elif not s:
            flush()
        else:
            para.append(s)
        i += 1
    flush()
    return title


def main():
    d = docx.Document(MANUAL)
    sid = {s.style_id: s.name for s in d.styles}
    body = d.element.body

    def is_h1(el):
        if el.tag != qn('w:p'):
            return False
        ppr = el.find(qn('w:pPr'))
        st = ppr.find(qn('w:pStyle')) if ppr is not None else None
        return st is not None and sid.get(st.get(qn('w:val'))) == 'Heading 1'

    for path in sys.argv[1:]:
        w = Writer(d)
        title = parse(path, w)
        els = list(body.iterchildren())
        heads = [i for i, el in enumerate(els) if is_h1(el) and ptext(el) == title]
        if len(heads) != 1:
            sys.exit('%s: %d headings "%s" in the manual' % (path, len(heads), title))
        h = heads[0]
        new = set(id(e) for e in w.out)
        old = []
        for el in els[h + 1:]:
            if is_h1(el) or el.tag == qn('w:sectPr') or id(el) in new:
                break
            if el.tag == qn('w:p') and el.find(qn('w:pPr')) is not None and \
                    el.find(qn('w:pPr')).find(qn('w:sectPr')) is not None:
                break                                   # (a section break stays where it is)
            old.append(el)
        for el in old:
            body.remove(el)
        anchor = els[h]
        for el in w.out:
            body.remove(el)
            anchor.addnext(el)
            anchor = el
        print('%s: "%s", %d elements in place of %d' % (os.path.basename(path), title, len(w.out), len(old)))
    d.save(MANUAL)


if __name__ == '__main__':
    main()
