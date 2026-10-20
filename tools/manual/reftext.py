"""reftext.py extract|build [NAME...] or reftext.py check COPY.docx [NAME...]: the reference sections of the manual
(their tables) as text in docs/manual/ref/NAME.txt, and back (docs/Manual_Plan.md, M4). The NAMEs (all of them when
none is given) are variables, options, commands, functions and obsolete.

The text is one entry (a row of the table) after another:

    === ADC OPEN freq, n_channels [, interrupt]      a row: its left cell's paragraphs, a line each
    =+= ADC FREQUENCY freq                           a row joined to the one above (no border between them:
                                                     the rows of one keyword)
    --- keep                                         the left cell ends; flags: P (Options: Permanent ticked),
                                                     keep (the row is not split across pages), wide (the wide
                                                     syntax column), tight and roomy (cell margins above and
                                                     below of 0 and 85 twips, not the table's)
    The ADC commands provide ...                     the right cell's paragraphs, a line each
    {}                                               an empty paragraph
    !TODO check the pins against the handler         a highlighted [MANUAL TODO] paragraph

A paragraph may start with its attributes in braces: b and a the spacing before and after, ln and lr the line
spacing and its rule, l, r, h and f the left, right, hanging and first line indents (twips), n the list (numbering
id) and v its level, j the alignment, t the tab stops (l720,r9000.dot: kind, position, leader), kn and kl keep with
next and keep lines together, z the size of an empty paragraph (half points).

Inline: **bold**, *italic*, __underlined__ (the applicability labels), `code` (Courier New), ^^superscript^^,
<<url>> or <<url|text>> a hyperlink, \\n a line break, a real TAB a tab, and U+2011 a non-breaking hyphen.
A backslash escapes the next character. A line "# ..." is a comment (the audit's notes), anywhere.

extract writes the table as it is; build replaces the table's rows by the file's entries (the header row of Options
is kept); check compares the table in a copy of the manual with the manual's, character by character with each
character's formatting, and their text forms."""
import copy, os, re, sys
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
MANUAL = os.path.join(ROOT, 'docs', 'manual', 'Colour_Maximite_2_User_Manual.docx')
TODO = '[MANUAL TODO] '
CODE_FONTS = ('Courier New', 'Lucida Console', 'Consolas', 'Courier')
PATTR = [('b', 'spacing', 'before'), ('a', 'spacing', 'after'), ('ln', 'spacing', 'line'),
         ('lr', 'spacing', 'lineRule'), ('l', 'ind', 'left'), ('r', 'ind', 'right'), ('h', 'ind', 'hanging'),
         ('f', 'ind', 'firstLine')]
TABKIND = {'left': 'l', 'right': 'r', 'center': 'c', 'decimal': 'd', 'bar': 'b', 'clear': 'x', 'num': 'n',
           'start': 'l', 'end': 'r'}
KINDTAB = {'l': 'left', 'r': 'right', 'c': 'center', 'd': 'decimal', 'b': 'bar', 'x': 'clear', 'n': 'num'}
MARK = {'c': '`', 'b': '**', 'i': '*', 'u': '__', 's': '^^'}
PLAIN = {'b': False, 'i': False, 'u': False, 'c': False, 's': False, 'link': None}
R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
SECTIONS = {'variables': 'Predefined Read Only Variables', 'options': 'Options', 'commands': 'Commands',
            'functions': 'Functions', 'obsolete': 'Obsolete Commands and Functions'}
REF = os.path.join(ROOT, 'docs', 'manual', 'ref')


def text_of(el):
    return ''.join(x.text or '' for x in el.iter(qn('w:t'))).strip()


def find_table(d, section):
    sid = {s.style_id: s.name for s in d.styles}
    els = list(d.element.body.iterchildren())
    for i, e in enumerate(els):
        if e.tag != qn('w:p'):
            continue
        ppr = e.find(qn('w:pPr'))
        st = ppr.find(qn('w:pStyle')) if ppr is not None else None
        if st is not None and sid.get(st.get(qn('w:val'))) == 'Heading 1' and text_of(e) == section:
            for e2 in els[i + 1:i + 20]:
                if e2.tag == qn('w:tbl'):
                    return e2
    sys.exit('no table under "%s"' % section)


def is_header(tr, section):
    tcs = tr.findall(qn('w:tc'))
    return section == 'Options' and len(tcs) > 1 and text_of(tcs[1]) == 'Permanent?'


# ---- reading a paragraph: its characters with their formatting ------------------------------------------------------

def on(parent, tag):
    el = parent.find(qn(tag)) if parent is not None else None
    return el is not None and el.get(qn('w:val')) not in ('0', 'false', 'none')


def run_fmt(r, link):
    rpr = r.find(qn('w:rPr'))
    f = rpr.find(qn('w:rFonts')) if rpr is not None else None
    font = (f.get(qn('w:ascii')) or f.get(qn('w:hAnsi')) or '') if f is not None else ''
    va = rpr.find(qn('w:vertAlign')) if rpr is not None else None
    return {'b': on(rpr, 'w:b'), 'i': on(rpr, 'w:i'), 'u': on(rpr, 'w:u'),
            'c': any(font.startswith(c) for c in CODE_FONTS),
            's': va is not None and va.get(qn('w:val')) == 'superscript', 'link': link}


def chars(p, rels):
    """the paragraph as (character, format) pairs: a tab is \\t, a line break \\n"""
    out = []

    def walk(el, link):
        for c in el:
            if c.tag == qn('w:r'):
                fmt = run_fmt(c, link)
                for x in c:
                    t = {qn('w:t'): None, qn('w:tab'): '\t', qn('w:br'): '\n', qn('w:cr'): '\n',
                         qn('w:noBreakHyphen'): '‑'}.get(x.tag, '')
                    for ch in ((x.text or '') if t is None else t):
                        out.append((ch, fmt))
            elif c.tag == qn('w:hyperlink'):
                rid = c.get('{%s}id' % R_NS)
                walk(c, rels[rid].target_ref if rid in rels else ('#' + (c.get(qn('w:anchor')) or '')))
            elif c.tag in (qn('w:smartTag'), qn('w:ins'), qn('w:fldSimple')):
                walk(c, link)
    walk(p, None)
    text = [ch for ch, _ in out]
    for k, (ch, fmt) in enumerate(out):                 # (bold or italic on a space does not show: a space is bold or
        if ch == ' ':                                   # italic when the characters on both sides of it are)
            a = next((j for j in range(k - 1, -1, -1) if text[j] != ' '), None)
            z = next((j for j in range(k + 1, len(out)) if text[j] != ' '), None)
            both = lambda m: a is not None and z is not None and out[a][1][m] and out[z][1][m]
            out[k] = (ch, dict(fmt, b=both('b'), i=both('i')))
    return out


def para_attrs(p):
    attrs = []
    ppr = p.find(qn('w:pPr'))
    if ppr is None:
        return attrs
    for tag, k in (('keepNext', 'kn'), ('keepLines', 'kl')):
        if on(ppr, 'w:' + tag):
            attrs.append(k + '=1')
    num = ppr.find(qn('w:numPr'))
    if num is not None and num.find(qn('w:numId')) is not None:
        attrs.append('n=%s' % num.find(qn('w:numId')).get(qn('w:val')))
        if num.find(qn('w:ilvl')) is not None:
            attrs.append('v=%s' % num.find(qn('w:ilvl')).get(qn('w:val')))
    tabs = ppr.find(qn('w:tabs'))
    if tabs is not None and len(tabs):
        ts = []
        for t in tabs:
            s = TABKIND.get(t.get(qn('w:val')), 'l') + t.get(qn('w:pos'))
            if t.get(qn('w:leader')) not in (None, 'none'):
                s += '.' + t.get(qn('w:leader'))
            ts.append(s)
        attrs.append('t=' + ','.join(ts))
    for k, tag, att in PATTR:
        el = ppr.find(qn('w:' + tag))
        if el is not None:
            v = el.get(qn('w:' + att))
            if v is None and tag == 'ind' and att in ('left', 'right'):     # (the newer names of the indents)
                v = el.get(qn('w:' + {'left': 'start', 'right': 'end'}[att]))
            if v is not None:
                attrs.append('%s=%s' % (k, v))
    jc = ppr.find(qn('w:jc'))
    if jc is not None:
        attrs.append('j=%s' % jc.get(qn('w:val')))
    rpr = ppr.find(qn('w:rPr'))
    sz = rpr.find(qn('w:sz')) if rpr is not None else None
    if sz is not None and not text_of(p):
        attrs.append('z=%s' % sz.get(qn('w:val')))
    return attrs


# ---- writing a paragraph as a line -----------------------------------------------------------------------------------

def esc(ch, prev, nxt):
    """ch as text; prev and nxt its neighbours with the same formatting (None at a change of formatting)"""
    if ch == '\n':
        return '\\n'
    if ch in '\\*`':
        return '\\' + ch
    if ch in '_^' and (prev == ch or nxt == ch or prev is None or nxt is None):
        return '\\' + ch                                 # (doubled, or at a mark: it could be read as one)
    if ch in '<>' and nxt == ch:
        return '\\' + ch
    return ch


def marked(cs):
    s, state = '', dict(PLAIN)
    for k, (ch, f) in enumerate(cs):
        for m in [m for m in 'suicb' if state[m] and not f[m]] + [m for m in 'bcius' if f[m] and not state[m]]:
            s += MARK[m]
            state[m] = f[m]
        prev = cs[k - 1][0] if k and cs[k - 1][1] == f else None
        nxt = cs[k + 1][0] if k + 1 < len(cs) and cs[k + 1][1] == f else None
        s += esc(ch, prev, nxt)
    return s + ''.join(MARK[m] for m in 'suicb' if state[m])


def para_line(p, rels):
    groups = []                                         # (link, its characters)
    for ch, f in chars(p, rels):
        if groups and groups[-1][0] == f['link']:
            groups[-1][1].append((ch, f))
        else:
            groups.append((f['link'], [(ch, f)]))
    s = ''
    for link, cs in groups:
        text = ''.join(c for c, _ in cs)
        if link is None:
            s += marked(cs)
        elif text == link and all(f == cs[0][1] and not any(f[m] for m in MARK) for _, f in cs):
            s += '<<' + link + '>>'
        else:
            s += '<<' + link + '|' + marked(cs) + '>>'
    attrs = para_attrs(p)
    if s.startswith(TODO.strip().replace('[', '\\[')) or s.startswith(TODO.strip()):
        return '!TODO ' + s.split(']', 1)[1].lstrip()
    if not attrs and s and (s[0] in '{#' or s[:4] in ('=== ', '=+= ', '---', '!TOD')):
        s = '\\' + s                                     # (a line that would read as something else)
    if not s and not attrs:
        return '{}'
    return ('{%s}' % ' '.join(attrs) if attrs else '') + s


def borders(tc):
    b = tc.find('.//' + qn('w:tcBorders'))
    get = lambda side: (b.find(qn('w:' + side)).get(qn('w:val'))
                        if b is not None and b.find(qn('w:' + side)) is not None else None)
    return get('top'), get('bottom')


def entry_lines(tr, section, grid, rels):
    tcs = tr.findall(qn('w:tc'))
    left, right = tcs[0], tcs[-1]
    lp = [para_line(p, rels) for p in left.findall(qn('w:p'))]
    out = [('=+= ' if borders(left)[0] == 'nil' else '=== ') + lp[0]] + lp[1:]
    flags = []
    if section == 'Options' and len(tcs) == 3 and tcs[1].find('.//' + qn('w:sym')) is not None:
        flags.append('P')
    trpr = tr.find(qn('w:trPr'))
    if trpr is not None and trpr.find(qn('w:cantSplit')) is not None:
        flags.append('keep')
    lw = int(left.find(qn('w:tcPr')).find(qn('w:tcW')).get(qn('w:w')))
    if len(tcs) == 2 and len(grid) == 3 and abs(lw - grid[0] - grid[1]) < 20:
        flags.append('wide')
    mar = left.find(qn('w:tcPr')).find(qn('w:tcMar'))
    top = mar.find(qn('w:top')) if mar is not None else None
    if top is not None:
        flags.append({'0': 'tight', '85': 'roomy'}.get(top.get(qn('w:w')), 'roomy'))
    out.append('---' + (' ' + ' '.join(flags) if flags else ''))
    return out + [para_line(p, rels) for p in right.findall(qn('w:p'))]


def table_lines(d, section):
    tbl = find_table(d, section)
    grid = [int(g.get(qn('w:w'))) for g in tbl.find(qn('w:tblGrid'))]
    out = []
    for tr in tbl.findall(qn('w:tr')):
        if not is_header(tr, section):
            out.append(entry_lines(tr, section, grid, d.part.rels))
    return out


def extract(section, path):
    d = docx.Document(MANUAL)
    entries = table_lines(d, section)
    lines = ['# %s: the reference entries (tools/manual/reftext.py)' % section, '']
    for e in entries:
        lines += e + ['']
    open(path, 'w', encoding='utf-8', newline='\n').write('\n'.join(lines))
    print('%s: %d entries written to %s' % (section, len(entries), path))


# ---- build -----------------------------------------------------------------------------------------------------------

def parse(path):
    entries, cur, side = [], None, None
    for n, raw in enumerate(open(path, encoding='utf-8').read().split('\n'), 1):
        if raw.startswith('# ') or raw == '#':
            continue
        if raw.startswith('=== ') or raw.startswith('=+= '):
            cur = {'joined': raw.startswith('=+='), 'left': [raw[4:]], 'right': [], 'flags': set(), 'line': n}
            entries.append(cur)
            side = 'left'
        elif cur is not None and side == 'left' and (raw == '---' or raw.startswith('--- ')):
            side = 'right'
            cur['flags'] = set(raw[3:].split())
            bad = cur['flags'] - {'P', 'keep', 'wide', 'tight', 'roomy'}
            if bad:
                sys.exit('%s:%d: unknown flag %s' % (path, n, ' '.join(sorted(bad))))
        elif raw != '':
            if cur is None:
                sys.exit('%s:%d: text before the first entry' % (path, n))
            cur[side].append(raw)
    return entries


TOKEN = re.compile(r'(\\.|<<|>>|\*\*|\^\^|__|`|\*)')


class Runs:
    """the inline marks of a line as the runs of paragraph p"""
    def __init__(self, d, p, hl=False):
        self.d, self.p, self.hl = d, p, hl
        self.state = dict(PLAIN)
        self.parent = p
        self.buf = []

    def rpr(self):
        rpr = OxmlElement('w:rPr')
        if self.parent is not self.p and any(s.name == 'Hyperlink' for s in self.d.styles):
            st = OxmlElement('w:rStyle'); st.set(qn('w:val'), self.d.styles['Hyperlink'].style_id); rpr.append(st)
        if self.state['c']:
            f = OxmlElement('w:rFonts')
            for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
                f.set(qn(a), 'Courier New')
            rpr.append(f)
        if self.state['b']:
            rpr.append(OxmlElement('w:b'))
        if self.state['i']:
            rpr.append(OxmlElement('w:i'))
        if self.hl:
            h = OxmlElement('w:highlight'); h.set(qn('w:val'), 'yellow'); rpr.append(h)
        if self.state['u']:
            u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rpr.append(u)
        if self.state['s']:
            v = OxmlElement('w:vertAlign'); v.set(qn('w:val'), 'superscript'); rpr.append(v)
        return rpr

    def flush(self):
        text, self.buf = ''.join(self.buf), []
        for piece in re.split('([\t\n‑])', text):
            if not piece:
                continue
            r = OxmlElement('w:r')
            r.append(self.rpr())
            if piece in ('\t', '\n', '‑'):
                r.append(OxmlElement({'\t': 'w:tab', '\n': 'w:br', '‑': 'w:noBreakHyphen'}[piece]))
            else:
                t = OxmlElement('w:t'); t.text = piece; t.set(qn('xml:space'), 'preserve'); r.append(t)
            self.parent.append(r)

    def add(self, text):
        tokens = TOKEN.split(text)
        k = 0
        while k < len(tokens):
            tok = tokens[k]
            if tok.startswith('\\') and len(tok) == 2:
                self.buf.append('\n' if tok[1] == 'n' else tok[1])
            elif tok == '<<' and k + 1 < len(tokens):  # <<url>> or <<url|text>>
                self.flush()
                url, bar, rest = tokens[k + 1].partition('|')
                h = OxmlElement('w:hyperlink')
                h.set('{%s}id' % R_NS, self.d.part.relate_to(url, RT.HYPERLINK, is_external=True))
                self.p.append(h)
                self.parent = h
                if bar:
                    tokens[k + 1] = rest
                else:
                    self.buf.append(url)
                    k += 1
            elif tok == '>>' and self.parent is not self.p:
                self.flush()
                self.parent = self.p
            elif tok in ('**', '*', '__', '`', '^^'):
                self.flush()
                m = {'**': 'b', '*': 'i', '__': 'u', '`': 'c', '^^': 's'}[tok]
                self.state[m] = not self.state[m]
            else:
                self.buf.append(tok)
            k += 1
        self.flush()


def make_para(d, line):
    p = OxmlElement('w:p')
    ppr = OxmlElement('w:pPr')
    attrs, text = {}, line
    m = re.match(r'^\{([^}]*)\}(.*)$', line)
    if m:
        attrs, text = dict(kv.split('=', 1) for kv in m.group(1).split()), m.group(2)
    for tag, k in (('keepNext', 'kn'), ('keepLines', 'kl')):
        if k in attrs:
            ppr.append(OxmlElement('w:' + tag))
    if 'n' in attrs:
        num = OxmlElement('w:numPr')
        il = OxmlElement('w:ilvl'); il.set(qn('w:val'), attrs.get('v', '0')); num.append(il)
        ni = OxmlElement('w:numId'); ni.set(qn('w:val'), attrs['n']); num.append(ni)
        ppr.append(num)
    if 't' in attrs:
        tabs = OxmlElement('w:tabs')
        for t in attrs['t'].split(','):
            el = OxmlElement('w:tab')
            pos, _, leader = t[1:].partition('.')
            el.set(qn('w:val'), KINDTAB[t[0]]); el.set(qn('w:pos'), pos)
            if leader:
                el.set(qn('w:leader'), leader)
            tabs.append(el)
        ppr.append(tabs)
    for tag in ('spacing', 'ind'):
        el = None
        for k, t, att in PATTR:
            if t == tag and k in attrs:
                if el is None:
                    el = OxmlElement('w:' + tag)
                el.set(qn('w:' + att), attrs[k])
        if el is not None:
            ppr.append(el)
    if 'j' in attrs:
        jc = OxmlElement('w:jc'); jc.set(qn('w:val'), attrs['j']); ppr.append(jc)
    if 'z' in attrs:
        rpr = OxmlElement('w:rPr')
        sz = OxmlElement('w:sz'); sz.set(qn('w:val'), attrs['z']); rpr.append(sz)
        ppr.append(rpr)
    p.append(ppr)
    if text.startswith('!TODO '):
        Runs(d, p, hl=True).add(TODO.replace('[', '\\[') + text[6:])
    else:
        Runs(d, p).add(text)
    return p


def set_nil(tc, side):
    tcpr = tc.find(qn('w:tcPr'))
    b = tcpr.find(qn('w:tcBorders'))
    if b is None:
        b = OxmlElement('w:tcBorders')
        later = [tcpr.find(qn('w:' + t)) for t in ('shd', 'noWrap', 'tcMar', 'textDirection', 'tcFitText',
                                                   'vAlign', 'hideMark')]
        later = [x for x in later if x is not None]
        if later:
            later[0].addprevious(b)
        else:
            tcpr.append(b)
    for el in b.findall(qn('w:' + side)):
        b.remove(el)
    el = OxmlElement('w:' + side)
    el.set(qn('w:val'), 'nil')
    order = ['top', 'left', 'start', 'bottom', 'right', 'end', 'insideH', 'insideV', 'tl2br', 'tr2bl']
    later = [c for c in b if c.tag.split('}')[1] in order[order.index(side) + 1:]]
    if later:
        later[0].addprevious(el)
    else:
        b.append(el)


def set_margins(tc, w):
    tcpr = tc.find(qn('w:tcPr'))
    m = OxmlElement('w:tcMar')
    for side in ('top', 'bottom'):
        el = OxmlElement('w:' + side); el.set(qn('w:w'), w); el.set(qn('w:type'), 'dxa'); m.append(el)
    later = [tcpr.find(qn('w:' + t)) for t in ('textDirection', 'tcFitText', 'vAlign', 'hideMark')]
    later = [x for x in later if x is not None]
    if later:
        later[0].addprevious(m)
    else:
        tcpr.append(m)


def build(section, path):
    d = docx.Document(MANUAL)
    tbl = find_table(d, section)
    grid = [g.get(qn('w:w')) for g in tbl.find(qn('w:tblGrid'))]
    body = [r for r in tbl.findall(qn('w:tr')) if not is_header(r, section)]
    template = copy.deepcopy(next(r for r in body if len(r.findall(qn('w:tc'))) == (3 if section == 'Options' else 2)))
    for tc in template.findall(qn('w:tc')):
        for p in tc.findall(qn('w:p')):
            tc.remove(p)
        tcpr = tc.find(qn('w:tcPr'))
        for b in tcpr.findall(qn('w:tcBorders')) + tcpr.findall(qn('w:tcMar')):
            tcpr.remove(b)
    for trpr in template.findall(qn('w:trPr')):
        template.remove(trpr)
    tick = None
    for r in body:                                      # (a ticked Permanent cell's paragraphs, to copy)
        tcs = r.findall(qn('w:tc'))
        if section == 'Options' and len(tcs) == 3 and tcs[1].find('.//' + qn('w:sym')) is not None:
            tick = [copy.deepcopy(p) for p in tcs[1].findall(qn('w:p'))]
            break
    for r in body:
        tbl.remove(r)
    entries = parse(path)
    prev = None
    for e in entries:
        tr = copy.deepcopy(template)
        tcs = tr.findall(qn('w:tc'))
        if 'keep' in e['flags']:
            trpr = OxmlElement('w:trPr'); trpr.append(OxmlElement('w:cantSplit')); tr.insert(0, trpr)
        if 'wide' in e['flags'] and len(tcs) == 2 and len(grid) == 3:
            for tc, w, span in ((tcs[0], int(grid[0]) + int(grid[1]), 2), (tcs[1], int(grid[2]), 1)):
                tcpr = tc.find(qn('w:tcPr'))
                tcpr.find(qn('w:tcW')).set(qn('w:w'), str(w))
                gs = tcpr.find(qn('w:gridSpan'))
                if span == 1 and gs is not None:
                    tcpr.remove(gs)
                elif span > 1:
                    if gs is None:
                        gs = OxmlElement('w:gridSpan'); tcpr.find(qn('w:tcW')).addnext(gs)
                    gs.set(qn('w:val'), str(span))
        for flag, w in (('tight', '0'), ('roomy', '85')):
            if flag in e['flags']:
                for tc in tcs:
                    set_margins(tc, w)
        for line in e['left']:
            tcs[0].append(make_para(d, line))
        for line in e['right']:
            tcs[-1].append(make_para(d, line))
        if len(tcs) == 3:
            for p in (tick if 'P' in e['flags'] and tick else [OxmlElement('w:p')]):
                tcs[1].append(copy.deepcopy(p))
        for tc in tcs:
            if not tc.findall(qn('w:p')):
                tc.append(OxmlElement('w:p'))
            if e['joined']:
                set_nil(tc, 'top')
        if e['joined'] and prev is not None:
            for tc in prev.findall(qn('w:tc')):
                set_nil(tc, 'bottom')
        tbl.append(tr)
        prev = tr
    d.save(MANUAL)
    print('%s: %d entries built from %s' % (section, len(entries), path))


def check(section, before):
    """the table in the copy BEFORE against the manual's: each cell's characters and formatting, and its text form"""
    forms, texts = [], []
    for f in (before, MANUAL):
        d = docx.Document(f)
        t = find_table(d, section)
        forms.append([[[(c, tuple(sorted(fm.items()))) for c, fm in chars(p, d.part.rels)] + [('|', ())]
                       for tc in tr.findall(qn('w:tc')) for p in tc.findall(qn('w:p'))]
                      for tr in t.findall(qn('w:tr'))])
        texts.append(table_lines(d, section))
    bad = 0
    for i, (ra, rb) in enumerate(zip(*forms)):
        sa, sb = [x for p in ra for x in p], [x for p in rb for x in p]
        if sa != sb:
            bad += 1
            if bad <= 4:
                k = next(j for j in range(min(len(sa), len(sb)) + 1) if j == min(len(sa), len(sb)) or sa[j] != sb[j])
                print('  row %d at %d: %r\n            %r' % (i, k, sa[k:k + 2], sb[k:k + 2]))
    tbad = sum(1 for x, y in zip(*texts) if x != y)
    for x, y in [(x, y) for x, y in zip(*texts) if x != y][:3]:
        print('  text form:', [l for l in x if l not in y][:2], '\n         ->', [l for l in y if l not in x][:2])
    print('%s: %d and %d rows; characters or formatting differ in %d, the text form in %d' %
          (section, len(forms[0]), len(forms[1]), bad + abs(len(forms[0]) - len(forms[1])), tbad))


if __name__ == '__main__':
    cmd, args = sys.argv[1], sys.argv[2:]
    before = args.pop(0) if cmd == 'check' else None
    for name in args or list(SECTIONS):
        path = os.path.join(REF, name + '.txt')
        if cmd == 'check':
            check(SECTIONS[name], before)
        else:
            {'extract': extract, 'build': build}[cmd](SECTIONS[name], path)
